"""État de la draft : ordre snake, keepers, choix saisis, joueurs libres."""

import csv
import logging

from scripts.names import name_key
from dataclasses import dataclass, field

log = logging.getLogger(__name__)


@dataclass
class Pick:
    overall: int          # 0 = premier choix de la draft
    round: int            # 1..rounds
    pick: int             # 1..teams (rang dans le tour)
    team: str
    name: str             # nom saisi
    player: object = None  # Player reconnu (None si inconnu)


@dataclass
class DraftState:
    order: list
    rounds: int
    my_team: str
    keepers: dict                         # manager -> [Player]
    picks: dict = field(default_factory=dict)   # overall -> Pick
    unknown: list = field(default_factory=list)  # noms non reconnus
    duplicates: list = field(default_factory=list)
    owners: list = None                   # ordre réel choix par choix (Yahoo), sinon snake

    @property
    def teams(self):
        return len(self.order)

    @property
    def total_picks(self):
        return self.teams * self.rounds

    def team_at(self, overall):
        if self.owners and overall < len(self.owners) and self.owners[overall]:
            return self.owners[overall]
        rnd, pos = divmod(overall, self.teams)
        return self.order[pos] if rnd % 2 == 0 else self.order[self.teams - 1 - pos]

    def slot(self, overall):
        rnd, pos = divmod(overall, self.teams)
        return rnd + 1, pos + 1

    @property
    def current(self):
        """Premier choix non encore saisi (None si draft terminée)."""
        for overall in range(self.total_picks):
            if overall not in self.picks:
                return overall
        return None

    def my_picks_from(self, start):
        return [o for o in range(start, self.total_picks) if self.team_at(o) == self.my_team and o not in self.picks]

    def rosters(self):
        """manager -> liste de Player (keepers + choix reconnus)."""
        result = {team: list(self.keepers.get(team, [])) for team in self.order}
        for pk in sorted(self.picks.values(), key=lambda p: p.overall):
            if pk.player is not None:
                result.setdefault(pk.team, []).append(pk.player)
        return result

    def taken(self):
        ids = {p.idx for players in self.keepers.values() for p in players}
        ids |= {pk.player.idx for pk in self.picks.values() if pk.player is not None}
        return ids


def build_state(league, pool, pick_rows, owners=None):
    """pick_rows : [(tour, choix_dans_le_tour, nom)] tels que saisis dans le classeur."""
    draft = league["draft"]
    order = list(draft.get("order") or [])
    teams = int(league["teams"])
    if len(order) != teams:
        raise ValueError(f"Onglet config : Draft | Order contient {len(order)} managers, {teams} attendus.")
    my_team = draft.get("my_team") or ""
    if my_team not in order:
        raise ValueError(f"Onglet config : My team '{my_team}' absent de Draft | Order.")
    if str(draft.get("type", "snake")).lower() != "snake":
        raise ValueError("Seule la draft snake est gérée pour l'instant.")

    unknown, seen, duplicates = [], set(), []
    keepers = {}
    for team, names in (draft.get("keepers") or {}).items():
        if team not in order:
            log.warning("Keepers : manager '%s' absent de draft.order, ignoré.", team)
            continue
        for name in names or []:
            player = pool.find(name)
            if player is None:
                unknown.append(f"{name} (keeper {team})")
                continue
            if player.idx in seen:   # keeper déclaré deux fois : seule la première déclaration compte
                duplicates.append(f"{player.name} (keeper {team}, déjà déclaré)")
                continue
            seen.add(player.idx)
            keepers.setdefault(team, []).append(player)

    state = DraftState(order=order, rounds=int(draft.get("rounds", 12)), my_team=my_team, keepers=keepers,
                       owners=list(owners) if owners else None)
    keeper_team = {p.idx: team for team, players in keepers.items() for p in players}
    owner = {p.idx: f"keeper de {team}" for team, players in keepers.items() for p in players}
    for rnd, pick, name in pick_rows:
        name = (name or "").strip()
        if not name:
            continue
        try:
            rnd, pick = int(float(rnd)), int(float(pick))
        except (TypeError, ValueError):
            continue
        if not (1 <= rnd <= state.rounds and 1 <= pick <= teams):
            continue
        overall = (rnd - 1) * teams + (pick - 1)
        player = pool.find(name)
        if player is None:
            unknown.append(f"{name} (tour {rnd}, choix {pick})")
        elif player.idx in keeper_team and keeper_team[player.idx] == state.team_at(overall):
            # keeper saisi dans les choix par son manager : compté comme ce choix, pas en double
            state.keepers[keeper_team.pop(player.idx)].remove(player)
        elif player.idx in seen:
            # déjà pris : le choix reste occupé (l'ordre snake avance) mais le joueur n'est pas compté
            # une 2e fois, sinon il gonflerait les stats d'une équipe
            duplicates.append(f"{player.name} (tour {rnd}, choix {pick} ; déjà {owner.get(player.idx, '?')})")
            player = None
        else:
            seen.add(player.idx)
            owner[player.idx] = f"choisi au tour {rnd}, choix {pick}"
        state.picks[overall] = Pick(overall, rnd, pick, state.team_at(overall), name, player)
    _fill_keeper_rounds(state, draft.get("keeper_rounds") or [])
    state.unknown, state.duplicates = unknown, duplicates
    for label in unknown:
        log.warning("Nom non reconnu : %s (ajoutez un alias dans l'onglet players)", label)
    for name in duplicates:
        log.warning("Joueur saisi deux fois, choix ignoré : %s", name)
    return state


def _fill_keeper_rounds(state, keeper_rounds):
    """Si les keepers occupent des tours de la draft (ex. [1, 2]), on les place d'office
    sur le choix de leur manager dans ces tours (sauf si le choix est déjà saisi)."""
    if not keeper_rounds:
        return
    for team, players in state.keepers.items():
        slots = [o for o in range(state.total_picks)
                 if state.team_at(o) == team and (o // state.teams) + 1 in keeper_rounds and o not in state.picks]
        for overall, player in zip(slots, list(players)):
            rnd, pick = state.slot(overall)
            state.picks[overall] = Pick(overall, rnd, pick, team, player.name, player)
            players.remove(player)


# --- Lecture des choix hors ligne (tests, mock drafts) ------------------------------------

def picks_from_xlsx(path, tab="draft_res"):
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)[tab]
    rows = []
    for values in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        rnd, pick, _key, name = (list(values) + [None] * 4)[:4]
        rows.append((rnd, pick, name))
    return rows


def picks_from_csv(path):
    """CSV avec colonnes round;pick;player (séparateur ; ou ,)."""
    with open(path, encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(2048)
        fh.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=";,")
        return [(r.get("round"), r.get("pick"), r.get("player")) for r in csv.DictReader(fh, dialect=dialect)]


def picks_from_assignments(league, pool, assignments):
    """[(manager, joueur)] (colonne Team Draft de l'onglet bdd) -> [(tour, choix, joueur)].

    L'onglet ne dit pas à quel choix chaque joueur a été pris : les joueurs d'un manager occupent
    ses choix dans l'ordre snake, ses keepers (Draft | Keepers) d'abord dans les tours keepers,
    les autres dans l'ordre des lignes. Le choix en cours est donc le premier choix d'un manager
    qui n'a pas encore autant de joueurs que de choix passés.
    """
    draft = league["draft"]
    order = list(draft.get("order") or [])
    teams, rounds = len(order), int(draft.get("rounds", 12))
    if not teams:
        return []
    keeper_rounds = set(draft.get("keeper_rounds") or [])
    keepers = {}
    for team, names in (draft.get("keepers") or {}).items():
        for name in names or []:
            player = pool.find(name)
            keepers.setdefault(team, set()).add(player.idx if player is not None else name_key(name))

    def team_at(overall):
        rnd, pos = divmod(overall, teams)
        return order[pos] if rnd % 2 == 0 else order[teams - 1 - pos]

    slots = {t: [o for o in range(teams * rounds) if team_at(o) == t] for t in order}
    rows, unknown, too_many = [], set(), set()
    by_team = {}
    for manager, name in assignments:
        if manager not in slots:
            unknown.add(manager)
            continue
        by_team.setdefault(manager, []).append(name)
    for team, names in by_team.items():
        def is_keeper(n):
            player = pool.find(n)
            return (player.idx if player is not None else name_key(n)) in keepers.get(team, set())
        keeper_slots = [o for o in slots[team] if o // teams + 1 in keeper_rounds]
        other_slots = [o for o in slots[team] if o // teams + 1 not in keeper_rounds]
        for name in [n for n in names if is_keeper(n)] + [n for n in names if not is_keeper(n)]:
            first, second = (keeper_slots, other_slots) if is_keeper(name) else (other_slots, keeper_slots)
            pool_slots = first or second
            if not pool_slots:
                too_many.add(team)
                continue
            overall = pool_slots.pop(0)
            rows.append((overall // teams + 1, overall % teams + 1, name))
    if unknown:
        log.warning("Team Draft : managers inconnus de Draft | Order ignorés : %s", ", ".join(sorted(unknown)))
    for team in sorted(too_many):
        log.warning("Team Draft : %s a plus de joueurs que de choix (%d), surplus ignoré.", team, rounds)
    return sorted(rows)


def picks_from_bdd_xlsx(path, league, pool, tab="bdd", gs=None):
    """Mock draft hors ligne : export Excel du classeur, colonne Team Draft de l'onglet bdd."""
    import openpyxl

    from scripts.bdd import draft_assignments

    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)[tab]
    rows = [["" if v is None else v for v in r] for r in ws.iter_rows(max_col=8, values_only=True)]
    return picks_from_assignments(league, pool, draft_assignments(rows, gs or {}))


def picks_from_sheet_rows(rows):
    """Lignes brutes d'une plage tour, choix, clé, joueur (ancien onglet draft_res)."""
    out = []
    for r in rows:
        r = list(r) + [""] * 4
        out.append((r[0], r[1], r[3]))
    return out
